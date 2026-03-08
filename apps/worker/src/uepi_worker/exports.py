"""
Export Generation - PDF and PPTX generation from analysis results
HealthForesight by DataMovers Branding
"""
from typing import Any, Optional
from datetime import datetime
from uuid import UUID
import json
from pathlib import Path
from io import BytesIO

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
    from reportlab.pdfgen import canvas
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from uepi_common.storage.factory import create_storage_client
from uepi_common.data.parquet_service import ParquetDataService

# HealthForesight Brand Colors (DataMovers DNA preserved)
HEALTHFORESIGHT_COLORS = {
    "primary": "#3B2F8F",  # DataMovers Indigo
    "secondary": "#1E2A44",  # Deep Navy
    "accent_start": "#2EC4C6",  # Teal start
    "accent_end": "#5EEAD4",  # Teal end
    "neutral_dark": "#0F172A",
    "neutral_mid": "#64748B",
    "neutral_light": "#E5E7EB",
    "positive": "#2E8B57",  # Muted Green
    "warning": "#E6A23C",  # Amber
    "risk": "#C04A4A",  # Muted Red
}


def hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    """Convert hex color to RGB tuple"""
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def hex_to_reportlab_color(hex_color: str) -> colors.Color:
    """Convert hex color to ReportLab Color"""
    r, g, b = hex_to_rgb(hex_color)
    return colors.Color(r/255.0, g/255.0, b/255.0)


def hex_to_pptx_color(hex_color: str) -> RGBColor:
    """Convert hex color to python-pptx RGBColor"""
    r, g, b = hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


class PDFGenerator:
    """Generate PDF exports from analysis results with HealthForesight branding"""
    
    def __init__(self, storage_client, bucket: str):
        self.storage_client = storage_client
        self.bucket = bucket
    
    def generate_audit_pack_pdf(
        self,
        tenant_id: UUID,
        analysis_id: UUID,
        analysis_results: dict[str, Any],
    ) -> bytes:
        """
        Generate audit pack PDF from analysis results with HealthForesight branding
        
        Args:
            tenant_id: Tenant ID
            analysis_id: Analysis ID
            analysis_results: Analysis results dictionary
        
        Returns:
            PDF bytes
        """
        if not REPORTLAB_AVAILABLE:
            # Fallback to placeholder if reportlab not available
            return self._generate_placeholder_pdf(analysis_id, analysis_results)
        
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=0.75*inch,
            leftMargin=0.75*inch,
            topMargin=1*inch,
            bottomMargin=0.75*inch,
        )
        
        # Create story (content elements)
        story = []
        
        # Define HealthForesight styles
        styles = getSampleStyleSheet()
        
        # Primary color styles
        primary_color = hex_to_reportlab_color(HEALTHFORESIGHT_COLORS["primary"])
        accent_color = hex_to_reportlab_color(HEALTHFORESIGHT_COLORS["accent_start"])
        neutral_dark = hex_to_reportlab_color(HEALTHFORESIGHT_COLORS["neutral_dark"])
        neutral_mid = hex_to_reportlab_color(HEALTHFORESIGHT_COLORS["neutral_mid"])
        
        # Title style
        title_style = ParagraphStyle(
            'HealthForesightTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=neutral_dark,
            spaceAfter=12,
            alignment=TA_LEFT,
            fontName='Helvetica-Bold',
        )
        
        # Heading style
        heading_style = ParagraphStyle(
            'HealthForesightHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=primary_color,
            spaceAfter=10,
            spaceBefore=12,
            fontName='Helvetica-Bold',
        )
        
        # Body style
        body_style = ParagraphStyle(
            'HealthForesightBody',
            parent=styles['BodyText'],
            fontSize=10,
            textColor=neutral_dark,
            spaceAfter=8,
            leading=14,
            fontName='Helvetica',
        )
        
        # Cover page
        story.append(Spacer(1, 2*inch))
        story.append(Paragraph("HealthForesight", title_style))
        story.append(Paragraph("by DataMovers", body_style))
        story.append(Spacer(1, 0.5*inch))
        story.append(Paragraph("Policy Impact Analysis", heading_style))
        story.append(Paragraph(f"Analysis ID: {analysis_id}", body_style))
        
        generated_date = datetime.utcnow().strftime("%B %d, %Y at %H:%M UTC")
        story.append(Paragraph(f"Generated: {generated_date}", body_style))
        story.append(PageBreak())
        
        # Executive Summary
        story.append(Paragraph("Executive Summary", title_style))
        impact_result = analysis_results.get('impact_result', {})
        effect_size = impact_result.get('effect_size', 0)
        percent_change = impact_result.get('percent_change', 0)
        
        summary_text = f"""
        This analysis examines the downstream impact of the policy implementation.
        We observed a change of {effect_size:.2f} units ({percent_change:+.1f}%) in the primary metric.
        """
        story.append(Paragraph(summary_text.strip(), body_style))
        story.append(Spacer(1, 0.2*inch))
        story.append(Paragraph("Why it matters:", heading_style))
        story.append(Paragraph(
            "Understanding policy impact before full rollout enables strategic adjustments that can improve outcomes while managing costs and access considerations.",
            body_style
        ))
        story.append(PageBreak())
        
        # Impact Results
        story.append(Paragraph("Impact Results", title_style))
        
        # Impact metrics table
        impact_data = [
            ['Metric', 'Pre-Period', 'Post-Period', 'Change', '% Change'],
        ]
        
        metrics = impact_result.get('metrics', {})
        treatment_pre = metrics.get('treatment_pre', {})
        treatment_post = metrics.get('treatment_post', {})
        
        for metric_key in ['utilization_per_1k', 'allowed_pmpm', 'paid_pmpm']:
            if metric_key in treatment_pre and metric_key in treatment_post:
                pre_val = treatment_pre[metric_key]
                post_val = treatment_post[metric_key]
                change = post_val - pre_val
                pct_change = (change / pre_val * 100) if pre_val != 0 else 0
                metric_name = metric_key.replace('_', ' ').title()
                impact_data.append([
                    metric_name,
                    f"{pre_val:.2f}",
                    f"{post_val:.2f}",
                    f"{change:+.2f}",
                    f"{pct_change:+.1f}%"
                ])
        
        impact_table = Table(impact_data)
        impact_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), primary_color),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        story.append(impact_table)
        story.append(PageBreak())
        
        # Trust Panel / Confidence & Limitations
        story.append(Paragraph("Confidence & Limitations", title_style))
        
        confidence_score = impact_result.get('confidence_score', {}).get('overall_score', 0)
        if isinstance(confidence_score, dict):
            confidence_score = confidence_score.get('overall_score', 0)
        if not isinstance(confidence_score, (int, float)):
            confidence_score = 0
        
        story.append(Paragraph(f"Confidence Score: {confidence_score:.0f}/100", heading_style))
        
        # Data sufficiency
        data_sufficiency = impact_result.get('data_sufficiency', {})
        if data_sufficiency:
            story.append(Paragraph("Data Coverage:", heading_style))
            coverage_text = f"""
            Sample Size: {data_sufficiency.get('actual_sample_size', 0):,} / {data_sufficiency.get('min_sample_size', 0):,} required
            Data Window: {data_sufficiency.get('data_window_months', 0)} months
            Coverage Score: {data_sufficiency.get('coverage_score', 0) * 100:.0f}%
            """
            story.append(Paragraph(coverage_text.strip(), body_style))
        
        # Limitations
        limitations = impact_result.get('limitations', [])
        if limitations:
            story.append(Paragraph("Limitations:", heading_style))
            for limitation in limitations:
                story.append(Paragraph(f"• {limitation}", body_style))
        
        story.append(PageBreak())
        
        # Methodology
        story.append(Paragraph("Methodology", title_style))
        methodology = impact_result.get('methodology', {})
        method_text = f"""
        Method: {methodology.get('method', 'Unknown')}
        Pre-Period: {methodology.get('pre_months', 6)} months
        Post-Period: {methodology.get('post_months', 6)} months
        Control Group: {'Yes' if methodology.get('has_control_group') else 'No'}
        """
        story.append(Paragraph(method_text.strip(), body_style))
        
        # Build PDF
        def add_footer(canvas_obj, doc):
            """Add footer to every page"""
            canvas_obj.saveState()
            canvas_obj.setFont('Helvetica', 8)
            canvas_obj.setFillColor(neutral_mid)
            canvas_obj.drawRightString(
                doc.pagesize[0] - 0.75*inch,
                0.5*inch,
                "Generated by HealthForesight by DataMovers"
            )
            canvas_obj.restoreState()
        
        doc.build(story, onFirstPage=add_footer, onLaterPages=add_footer)
        
        return buffer.getvalue()
    
    def _generate_placeholder_pdf(self, analysis_id: UUID, analysis_results: dict[str, Any]) -> bytes:
        """Generate placeholder PDF when reportlab is not available"""
        generated_date = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
        
        pdf_content = f"""
HealthForesight by DataMovers
AUDIT PACK - Analysis {analysis_id}

Generated: {generated_date}

Analysis Results:
{json.dumps(analysis_results, indent=2, default=str)}

[Placeholder PDF - reportlab not available]
Please install reportlab to generate full PDF exports:
pip install reportlab pillow
"""
        return pdf_content.encode('utf-8')


class PPTXGenerator:
    """Generate PPTX exports from analysis results with HealthForesight branding"""
    
    def __init__(self, storage_client, bucket: str):
        self.storage_client = storage_client
        self.bucket = bucket
    
    def generate_presentation_pptx(
        self,
        tenant_id: UUID,
        analysis_id: UUID,
        analysis_results: dict[str, Any],
    ) -> bytes:
        """
        Generate presentation PPTX from analysis results with HealthForesight branding
        
        Args:
            tenant_id: Tenant ID
            analysis_id: Analysis ID
            analysis_results: Analysis results dictionary
        
        Returns:
            PPTX bytes
        """
        if not PPTX_AVAILABLE:
            # Fallback to placeholder if python-pptx not available
            return self._generate_placeholder_pptx(analysis_id, analysis_results)
        
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        
        # HealthForesight colors
        primary_color = hex_to_pptx_color(HEALTHFORESIGHT_COLORS["primary"])
        accent_color = hex_to_pptx_color(HEALTHFORESIGHT_COLORS["accent_start"])
        neutral_dark = hex_to_pptx_color(HEALTHFORESIGHT_COLORS["neutral_dark"])
        neutral_mid = hex_to_pptx_color(HEALTHFORESIGHT_COLORS["neutral_mid"])
        
        # Title slide
        title_slide_layout = prs.slide_layouts[0]  # Title slide layout
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "HealthForesight"
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = neutral_dark
        
        subtitle.text = f"Policy Impact Analysis\nby DataMovers\n\nAnalysis ID: {analysis_id}"
        subtitle.text_frame.paragraphs[0].font.size = Pt(18)
        subtitle.text_frame.paragraphs[0].font.color.rgb = neutral_mid
        
        # Executive Summary slide
        bullet_slide_layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = "Executive Summary"
        title_shape.text_frame.paragraphs[0].font.size = Pt(32)
        title_shape.text_frame.paragraphs[0].font.color.rgb = primary_color
        title_shape.text_frame.paragraphs[0].font.bold = True
        
        tf = body_shape.text_frame
        tf.text = "What changed?"
        
        impact_result = analysis_results.get('impact_result', {})
        effect_size = impact_result.get('effect_size', 0)
        percent_change = impact_result.get('percent_change', 0)
        
        p = tf.add_paragraph()
        p.text = f"We observed a change of {effect_size:.2f} units ({percent_change:+.1f}%) in the primary metric."
        p.level = 1
        p.font.size = Pt(14)
        p.font.color.rgb = neutral_dark
        
        p = tf.add_paragraph()
        p.text = "Why it matters"
        p.level = 0
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = primary_color
        
        p = tf.add_paragraph()
        p.text = "Understanding policy impact before full rollout enables strategic adjustments."
        p.level = 1
        p.font.size = Pt(14)
        p.font.color.rgb = neutral_dark
        
        # Impact Results slide
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = "Impact Results"
        title_shape.text_frame.paragraphs[0].font.size = Pt(32)
        title_shape.text_frame.paragraphs[0].font.color.rgb = primary_color
        title_shape.text_frame.paragraphs[0].font.bold = True
        
        tf = body_shape.text_frame
        metrics = impact_result.get('metrics', {})
        treatment_pre = metrics.get('treatment_pre', {})
        treatment_post = metrics.get('treatment_post', {})
        
        for metric_key in ['utilization_per_1k', 'allowed_pmpm', 'paid_pmpm']:
            if metric_key in treatment_pre and metric_key in treatment_post:
                pre_val = treatment_pre[metric_key]
                post_val = treatment_post[metric_key]
                change = post_val - pre_val
                pct_change = (change / pre_val * 100) if pre_val != 0 else 0
                metric_name = metric_key.replace('_', ' ').title()
                p = tf.add_paragraph()
                p.text = f"{metric_name}: {pre_val:.2f} → {post_val:.2f} ({pct_change:+.1f}%)"
                p.font.size = Pt(14)
                p.font.color.rgb = neutral_dark
        
        # Confidence & Limitations slide
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = "Confidence & Limitations"
        title_shape.text_frame.paragraphs[0].font.size = Pt(32)
        title_shape.text_frame.paragraphs[0].font.color.rgb = primary_color
        title_shape.text_frame.paragraphs[0].font.bold = True
        
        tf = body_shape.text_frame
        
        confidence_score = impact_result.get('confidence_score', {}).get('overall_score', 0)
        if isinstance(confidence_score, dict):
            confidence_score = confidence_score.get('overall_score', 0)
        if not isinstance(confidence_score, (int, float)):
            confidence_score = 0
        
        p = tf.add_paragraph()
        p.text = f"Confidence Score: {confidence_score:.0f}/100"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = primary_color
        
        limitations = impact_result.get('limitations', [])
        if limitations:
            p = tf.add_paragraph()
            p.text = "Limitations:"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = neutral_dark
            
            for limitation in limitations[:5]:  # Limit to 5 for readability
                p = tf.add_paragraph()
                p.text = f"• {limitation}"
                p.level = 1
                p.font.size = Pt(12)
                p.font.color.rgb = neutral_mid
        
        # Add footer to all slides (except title)
        generated_date = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
        for idx, slide in enumerate(prs.slides):
            if idx > 0:  # Skip title slide
                # Add footer text box at bottom
                left = Inches(0.5)
                top = Inches(6.8)
                width = Inches(9)
                height = Inches(0.3)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.text = f"Generated by HealthForesight by DataMovers | {generated_date}"
                tf.paragraphs[0].font.size = Pt(8)
                tf.paragraphs[0].font.color.rgb = neutral_mid
                tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        
        # Save to bytes
        buffer = BytesIO()
        prs.save(buffer)
        return buffer.getvalue()
    
    def _generate_placeholder_pptx(self, analysis_id: UUID, analysis_results: dict[str, Any]) -> bytes:
        """Generate placeholder PPTX when python-pptx is not available"""
        generated_date = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
        
        pptx_content = f"""
HealthForesight by DataMovers
PRESENTATION - Analysis {analysis_id}

Generated: {generated_date}

Analysis Results:
{json.dumps(analysis_results, indent=2, default=str)}

[Placeholder PPTX - python-pptx not available]
Please install python-pptx to generate full PPTX exports:
pip install python-pptx
"""
        return pptx_content.encode('utf-8')
